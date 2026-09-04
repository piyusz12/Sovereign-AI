"""
Sovereign AI Workbench — Deliverable Validator

Ensures that generated deliverables (DOCX, XLSX, PPTX, ZIP) are structurally
sound, not corrupt, and contain expected minimum content before releasing them.
"""

import logging
from pathlib import Path
import zipfile
from dataclasses import dataclass

logger = logging.getLogger("sovereign.generators.validator")

@dataclass
class ValidationResult:
    is_valid: bool
    error: str = ""


class DeliverableValidator:
    """Validates the structure and integrity of generated files."""

    @staticmethod
    def validate_docx(filepath: str) -> ValidationResult:
        """Verify that the DOCX opens and is not corrupt."""
        try:
            from docx import Document
            doc = Document(filepath)
            # Ensure it has at least one paragraph
            if len(doc.paragraphs) == 0:
                return ValidationResult(False, "Generated DOCX is empty.")
            return ValidationResult(True)
        except Exception as e:
            logger.error(f"DOCX Validation failed for {filepath}: {e}")
            return ValidationResult(False, f"Corrupted DOCX file: {e}")

    @staticmethod
    def validate_xlsx(filepath: str) -> ValidationResult:
        """Verify that the XLSX opens via openpyxl and has sheets."""
        try:
            import openpyxl
            wb = openpyxl.load_workbook(filepath, read_only=True)
            if not wb.sheetnames:
                return ValidationResult(False, "Generated XLSX has no sheets.")
            wb.close()
            return ValidationResult(True)
        except Exception as e:
            logger.error(f"XLSX Validation failed for {filepath}: {e}")
            return ValidationResult(False, f"Corrupted XLSX file: {e}")

    @staticmethod
    def validate_pptx(filepath: str) -> ValidationResult:
        """Verify that the PPTX opens and has slides."""
        try:
            from pptx import Presentation
            prs = Presentation(filepath)
            if len(prs.slides) == 0:
                return ValidationResult(False, "Generated PPTX has no slides.")
            return ValidationResult(True)
        except Exception as e:
            logger.error(f"PPTX Validation failed for {filepath}: {e}")
            return ValidationResult(False, f"Corrupted PPTX file: {e}")

    @staticmethod
    def validate_code_package(filepath: str, required_files: list[str] = None) -> ValidationResult:
        """Verify that the ZIP file opens and contains required files."""
        if required_files is None:
            required_files = ["main.py"]
            
        try:
            with zipfile.ZipFile(filepath, 'r') as zf:
                files = zf.namelist()
                missing = [f for f in required_files if f not in files]
                if missing:
                    return ValidationResult(False, f"Code package missing required files: {missing}")
            return ValidationResult(True)
        except Exception as e:
            logger.error(f"ZIP Validation failed for {filepath}: {e}")
            return ValidationResult(False, f"Corrupted ZIP package: {e}")


# Global instance
validator = DeliverableValidator()
